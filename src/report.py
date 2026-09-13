"""report.

Turn a stream of MatchResults into one spreadsheet that shows EVERY row -
successes and failures together - which is what you asked for. We write a
fresh file and never touch the original source of truth.

Each verdict is just *read* here; none is *decided* here. That separation is
why the same MatchResults could instead feed a loader with no rewrite.
"""
from __future__ import annotations

from typing import Iterable

import openpyxl
from openpyxl.styles import Font, PatternFill

from .normalise import MatchLevel
from .matcher import MatchResult

# one colour per verdict, best -> worst
_FILL = {
    MatchLevel.EXACT:  "C6EFCE",   # green  - clean match
    MatchLevel.SAFE:   "FFF2CC",   # amber  - fine, but note it
    MatchLevel.LOOSE:  "FCE4D6",   # orange - human must review
    MatchLevel.NONE:   "F8CBAD",   # red    - missing
}
_HEADERS = ["File Path", "disk_path", "status", "cause"]


def _row(r: MatchResult) -> list:
    return [
        r.expected,
        r.resolved or "",
        r.level.name,
        r.cause,
    ]


def write_report(results: Iterable[MatchResult], out_path: str) -> dict:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Match Report"

    ws.append(_HEADERS)
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True)

    tally: dict[str, int] = {}
    for r in results:
        ws.append(_row(r))
        tally[r.level.name] = tally.get(r.level.name, 0) + 1
        fill = PatternFill("solid", fgColor=_FILL[r.level])
        for c in ws[ws.max_row]:
            c.fill = fill
            c.font = Font(name="Arial")

    # widen columns so the report is readable on open
    for col, width in zip("ABCD", (34, 34, 9, 42)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"           # keep the header visible when scrolling

    wb.save(out_path)
    return tally