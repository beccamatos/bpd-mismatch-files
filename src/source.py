
"""
Where the expected paths come from is a separate concern from how we match
them. Keeping this here means the matcher never imports openpyxl, and swapping
Excel for a CSV or a plain list later touches only this file.
"""
from __future__ import annotations
 
import openpyxl
 
 
def read_xlsx(path: str, sheet: str | None = None,
              column: str = "File Path") -> list[str]:
    """Return the expected paths from one column of a spreadsheet.
 
    Finds the column by its header text (not a fixed letter) so the sheet can
    be rearranged without breaking us. Skips blank cells.
    """
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[sheet] if sheet else wb.active
 
    header = [ (c.value or "") for c in next(ws.iter_rows(min_row=1, max_row=1)) ]
    try:
        col_idx = header.index(column)          # 0-based position of our column
    except ValueError:
        raise ValueError(f"column {column!r} not found; headers are {header}")
 
    paths: list[str] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        value = row[col_idx] if col_idx < len(row) else None
        if value is not None and str(value).strip() != "":
            paths.append(str(value))
    return paths
 